"""The stats spreadsheet: every wheel result (governments, ideologies, history, traits,
resources…) and what it gives, as an Excel workbook you can tune, and back again.

    python scripts/stats_sheet.py export [file.xlsx]   # data file -> workbook (default tuning/1500-stats.xlsx)
    python scripts/stats_sheet.py import [file.xlsx]   # workbook -> public/data/eras/1500-stats.json

Export reads public/data/eras/1500-stats.json (the numbers) and public/data/eras/1500.json
(names, descriptions, chart positions). Import writes the numbers you changed back into the
stats file and leaves everything else in it as it was. Needs Python 3 with openpyxl.
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
STATS = ROOT / "public/data/eras/1500-stats.json"
WHEELS = ROOT / "public/data/eras/1500.json"
DEFAULT_XLSX = ROOT / "tuning/1500-stats.xlsx"

# The columns: (factor id, heading, "pct" or "pts"), in groups.
GROUPS = [
    ("Military", [("recruitRate", "Recruit rate", "pct"), ("armyShare", "Kept under arms", "pct"), ("soldierCost", "Cost per soldier", "pct"),
                  ("armyBudget", "Army budget", "pct"), ("armyQuality", "Army quality", "pts"), ("aggression", "Aggression", "pts")]),
    ("Economy", [("taxRate", "Tax rate", "pct"), ("adminCost", "Court & officials", "pct"), ("treasuryYears", "Savings", "pct")]),
    ("People", [("growthRate", "Birth rate", "pct"), ("food", "Farming", "pct"), ("stability", "Stability", "pts"),
                ("popularity", "Popularity", "pts"), ("development", "Development", "pts")]),
    ("Land & power", [("defence", "Defence", "pts"), ("trade", "Trade", "pts"), ("naval", "Naval reach", "pts"), ("cohesion", "Cohesion", "pts")]),
]
FACTORS = [f for _, cols in GROUPS for f in cols]
KIND = {fid: kind for fid, _, kind in FACTORS}

# The tabs of wheel results: (tab name, wheel id, what the rows are).
TABS = [
    ("Governments", "government", "Each government type"),
    ("Ideologies", "ideology", "Each ideology"),
    ("History", "history", "Each history result"),
    ("Bonus traits", "bonus", "Each result of the bonus spin"),
    ("Golden traits", "golden", "Each result of the golden spin"),
    ("Wealth", "wealth", "Each wealth level"),
    ("Resources", "resource", "Each bonus resource"),
    ("Stability words", "stability", "Each word on the Stability wheel (its number: 50 + this)"),
    ("Popularity words", "acceptance", "Each word on the Popularity wheel (its number: 50 + this)"),
]

FONT = "Arial"
GOLD = "B8913A"
INK = "2B2418"
HEAD_FILL = PatternFill("solid", fgColor="2E271D")
GROUP_FILLS = {"Military": "F6DCD6", "Economy": "F8EBC8", "People": "DCEFD6", "Land & power": "DDE8F2"}
EDIT_FILL = PatternFill("solid", fgColor="FFFBEA")
INFO_FILL = PatternFill("solid", fgColor="EFEFEF")
THIN = Side(style="thin", color="D8CDB4")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PCT_FORMAT = '[Color10]+0%;[Red]-0%;0%'
PTS_FORMAT = '[Color10]+0;[Red]-0;0'


def parse(text):
    """'+30%' -> 0.30, '-5' -> -5.0, '×0' -> None (not a sheet value)."""
    t = str(text).strip().replace("−", "-")
    if t[:1] in "×x*":
        return None
    if t.endswith("%"):
        return float(t[:-1]) / 100
    return float(t)


def fmt(value, kind):
    """0.30 -> '+30%', -5 -> '-5'."""
    v = value * 100 if kind == "pct" else value
    v = round(v, 1)
    text = str(int(v)) if v == int(v) else str(v)
    sign = "+" if v >= 0 else ""
    return f"{sign}{text}{'%' if kind == 'pct' else ''}"


def wheel_slices(wheels, wheel_id):
    """The wheel's slices in order, with rare results (on the rare slice's own wheel) after."""
    wheel = next(w for w in wheels["wheels"] if w["id"] == wheel_id)
    out = []
    for s in wheel["slices"]:
        if s.get("then"):
            for r in s["then"]["slices"]:
                out.append({**r, "rare": True})
        else:
            out.append(s)
    return [s for s in out if s.get("id") != "rare"]


def cell(ws, row, col, value=None, *, bold=False, color=INK, size=10, fill=None, fmt_=None, align=None, italic=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
    if fill:
        c.fill = fill
    if fmt_:
        c.number_format = fmt_
    if align:
        c.alignment = align
    return c


def export(path):
    stats = json.loads(STATS.read_text(encoding="utf8"))
    wheels = json.loads(WHEELS.read_text(encoding="utf8"))
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    write_readme(readme, stats)
    for title, wheel_id, what in TABS:
        write_tab(wb.create_sheet(title), wheel_id, what, stats, wheels)
    write_effects(wb.create_sheet("Stability effects"), stats)
    write_bases(wb.create_sheet("Starting values"), stats)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"Wrote {path}")


def write_readme(ws, stats):
    lines = [
        ("Little Dictators · 1500 stats", "title"),
        ("What every government, ideology, history, trait, wealth level and resource gives a nation.", None),
        ("", None),
        ("How to read it", "head"),
        ("Each row is one wheel result; each column is one of the numbers the game works out stats from.", None),
        ("Green = raises it, red = lowers it, empty = no effect.", None),
        ("% columns change the number by a share of its starting value: +30% recruit rate turns 2.0% into 2.6%. Several % add up first.", None),
        ("pts columns add points to a 0–100 score: +20 aggression, -10 stability.", None),
        ("Grey columns (id, chart position, description) are for information only.", None),
        ("", None),
        ("How to change it", "head"),
        ("Edit the cream cells: type 30% (or 0.3) for +30%, -15% for -15%, 20 for +20 points. Clear a cell to remove an effect.", None),
        ("Don't change the ids in column A or the factor names in row 3; the game finds rows and columns by them.", None),
        ("'Stability effects' says what stability and popularity do to everything (at 0, at 100; nothing at 50).", None),
        ("'Starting values' has where each number starts, and its lowest and highest.", None),
        ("When you're done, save the file and send it back; the game reads it with: python scripts/stats_sheet.py import", None),
        ("", None),
        ("Example", "head"),
        ("Militarist (Ideologies tab): Recruit rate +30%, Kept under arms +30%, Army budget +20%, Army quality +10, Aggression +20.", None),
        ("", None),
        ("Where the numbers start", "head"),
    ]
    for i, (text, style) in enumerate(lines, start=1):
        if style == "title":
            cell(ws, i, 1, text, bold=True, size=16, color=GOLD)
        elif style == "head":
            cell(ws, i, 1, text, bold=True, size=11)
        else:
            cell(ws, i, 1, text)
    row = len(lines) + 1
    for fid, name, kind in FACTORS:
        f = stats["factors"][fid]
        base = f["base"]
        shown = f"{base * 100:g}%" if f["format"] == "percent" else f"${base:,.0f}" if f["format"] == "money" else f"{base:g}"
        cell(ws, row, 1, f"{name}: starts at {shown}. {f.get('about', '')}".strip())
        row += 1
    ws.column_dimensions["A"].width = 130


def write_tab(ws, wheel_id, what, stats, wheels):
    mods = stats["results"].get(wheel_id, {})
    slices = wheel_slices(wheels, wheel_id)
    info = ["id", "Name", "Kind", "Chart x", "Description"]
    first = len(info) + 1
    # Row 1: group bands; row 2: headings; row 3: factor ids (the import reads these).
    cell(ws, 1, 1, what, bold=True, size=11, color=GOLD)
    col = first
    for group, cols in GROUPS:
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(cols) - 1)
        cell(ws, 1, col, group, bold=True, fill=PatternFill("solid", fgColor=GROUP_FILLS[group]), align=Alignment(horizontal="center"))
        col += len(cols)
    for i, h in enumerate(info, start=1):
        cell(ws, 2, i, h, bold=True, color="F1E8D4", fill=HEAD_FILL)
        cell(ws, 3, i, "", fill=HEAD_FILL)
    for j, (fid, name, kind) in enumerate(FACTORS):
        c = cell(ws, 2, first + j, f"{name} ({'%' if kind == 'pct' else 'pts'})", bold=True, color="F1E8D4", fill=HEAD_FILL,
                 align=Alignment(horizontal="center", wrap_text=True))
        about = stats["factors"][fid].get("about")
        if about:
            c.comment = Comment(about, "Little Dictators")
        cell(ws, 3, first + j, fid, size=8, color="9A8F78", fill=HEAD_FILL, align=Alignment(horizontal="center"))
    ws.row_dimensions[2].height = 30
    for r, s in enumerate(slices, start=4):
        kind = "Rare" if s.get("rare") else "Golden" if s.get("golden") else ""
        chart = (s.get("chart") or [None, None])[0]
        for i, v in enumerate([s["id"], s["label"], kind, chart, s.get("blurb", "")], start=1):
            c = cell(ws, r, i, v, fill=INFO_FILL if i != 2 else None, bold=i == 2, color="C0283A" if kind == "Rare" and i == 3 else INK)
            c.border = BOX
            if i == 4:
                c.number_format = "+0.0;-0.0;0"
        row_mods = mods.get(s["id"], {})
        for j, (fid, _, kindf) in enumerate(FACTORS):
            raw = row_mods.get(fid)
            value = parse(raw) if raw is not None else None
            c = cell(ws, r, first + j, value, fill=EDIT_FILL, fmt_=PCT_FORMAT if kindf == "pct" else PTS_FORMAT, align=Alignment(horizontal="center"))
            c.border = BOX
    widths = [20, 26, 8, 8, 44] + [11] * len(FACTORS)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=4, column=3)
    ws.auto_filter.ref = f"A3:{get_column_letter(len(widths))}{3 + len(slices)}"


def write_effects(ws, stats):
    cell(ws, 1, 1, "What stability and popularity do to everything else: the effect at 0 and at 100 (nothing at 50, half at 25 or 75).", bold=True, color=GOLD)
    for i, h in enumerate(["Score", "Changes", "factor id", "At 0", "At 100"], start=1):
        cell(ws, 2, i, h, bold=True, color="F1E8D4", fill=HEAD_FILL)
    names = {fid: name for fid, name, _ in FACTORS}
    r = 3
    for score in ("stability", "popularity"):
        for fid, pair in stats["scoreEffects"].get(score, {}).items():
            if fid == "about":
                continue
            kind = "pct" if str(pair[0]).strip().endswith("%") else "pts"
            cell(ws, r, 1, score.capitalize(), fill=INFO_FILL)
            cell(ws, r, 2, f"{names.get(fid, fid)} ({'%' if kind == 'pct' else 'pts'})", fill=INFO_FILL)
            cell(ws, r, 3, fid, size=8, color="9A8F78", fill=INFO_FILL)
            for k in (0, 1):
                cell(ws, r, 4 + k, parse(pair[k]), fill=EDIT_FILL, fmt_=PCT_FORMAT if kind == "pct" else PTS_FORMAT, align=Alignment(horizontal="center"))
            r += 1
    for i, w in enumerate([14, 26, 14, 10, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"


def write_bases(ws, stats):
    cell(ws, 1, 1, "Where each number starts before any result changes it, and the lowest and highest it can go.", bold=True, color=GOLD)
    for i, h in enumerate(["Number", "factor id", "Start", "Lowest", "Highest", "Unit", "What it is"], start=1):
        cell(ws, 2, i, h, bold=True, color="F1E8D4", fill=HEAD_FILL)
    units = {"percent": "share (0.02 = 2%)", "money": "$ a year", "score": "points (0–100)", "times": "× (1 = normal)", "years": "years of income"}
    for r, (fid, name, _) in enumerate(FACTORS, start=3):
        f = stats["factors"][fid]
        cell(ws, r, 1, name, fill=INFO_FILL, bold=True)
        cell(ws, r, 2, fid, size=8, color="9A8F78", fill=INFO_FILL)
        for k, key in enumerate(("base", "min", "max")):
            cell(ws, r, 3 + k, f[key], fill=EDIT_FILL, align=Alignment(horizontal="center"))
        cell(ws, r, 6, units.get(f["format"], f["format"]), fill=INFO_FILL)
        cell(ws, r, 7, f.get("about", ""), fill=INFO_FILL)
    for i, w in enumerate([22, 14, 10, 10, 10, 18, 80], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"


def read_value(v):
    """A cell as a number: 0.3, '30%', '+30%', '-5', None/'' -> None."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace("−", "-").replace(",", ".")
    if t.endswith("%"):
        return float(t[:-1]) / 100
    return float(t)


def import_(path):
    stats = json.loads(STATS.read_text(encoding="utf8"))
    text = STATS.read_text(encoding="utf8")
    wb = load_workbook(path, data_only=True)
    changes = 0
    problems = []
    for title, wheel_id, _ in TABS:
        if title not in wb.sheetnames:
            problems.append(f"The tab '{title}' is missing.")
            continue
        ws = wb[title]
        ids = [ws.cell(row=3, column=c).value for c in range(1, ws.max_column + 1)]
        old = stats["results"].get(wheel_id, {})
        for r in range(4, ws.max_row + 1):
            sid = ws.cell(row=r, column=1).value
            if not sid:
                continue
            mods = {}
            for c, fid in enumerate(ids, start=1):
                if fid not in KIND:
                    continue
                try:
                    v = read_value(ws.cell(row=r, column=c).value)
                except ValueError:
                    problems.append(f"{title}, {sid}, {fid}: '{ws.cell(row=r, column=c).value}' isn't a number.")
                    continue
                if v is None or v == 0:
                    continue
                if KIND[fid] == "pct" and abs(v) > 5:
                    v = v / 100  # typed 30 for 30%
                mods[fid] = fmt(v, KIND[fid])
            prev = old.get(sid, {})
            # Keep the order the file had, then any new effects.
            ordered = {k: mods[k] for k in prev if k in mods}
            ordered.update({k: v for k, v in mods.items() if k not in ordered})
            if ordered == {k: v for k, v in prev.items() if k in KIND} and set(prev) <= set(KIND):
                continue
            changes += 1
            text = set_line(text, ["results", wheel_id], sid, one_line(ordered) if ordered else None)
    if "Stability effects" in wb.sheetnames:
        ws = wb["Stability effects"]
        for r in range(3, ws.max_row + 1):
            score, fid = ws.cell(row=r, column=1).value, ws.cell(row=r, column=3).value
            if not score or fid not in KIND:
                continue
            key = score.lower()
            old_pair = stats["scoreEffects"][key].get(fid)
            kind = "pct" if old_pair and str(old_pair[0]).strip().endswith("%") else KIND[fid]
            pair = []
            for c in (4, 5):
                v = read_value(ws.cell(row=r, column=c).value) or 0.0
                if kind == "pct" and abs(v) > 5:
                    v = v / 100
                pair.append(fmt(v, kind))
            if pair != old_pair:
                changes += 1
                text = set_line(text, ["scoreEffects", key], fid, one_line(pair))
    if "Starting values" in wb.sheetnames:
        ws = wb["Starting values"]
        for r in range(3, ws.max_row + 1):
            fid = ws.cell(row=r, column=2).value
            if fid not in stats["factors"]:
                continue
            spec = dict(stats["factors"][fid])
            for c, key in ((3, "base"), (4, "min"), (5, "max")):
                v = read_value(ws.cell(row=r, column=c).value)
                if v is not None:
                    spec[key] = int(v) if v == int(v) else round(v, 6)
            if spec != stats["factors"][fid]:
                changes += 1
                text = set_line(text, ["factors"], fid, one_line(spec))
    if problems:
        print("Not imported, fix these first:\n  " + "\n  ".join(problems))
        sys.exit(1)
    json.loads(text)  # still valid JSON
    STATS.write_text(text, encoding="utf8")
    print(f"Updated {STATS.relative_to(ROOT)}: {changes} rows changed.")


def one_line(v):
    """A value on one line, like the hand-written file: { "a": "+5", "b": "-10%" }."""
    if isinstance(v, dict):
        return "{ " + ", ".join(f"{json.dumps(k, ensure_ascii=False)}: {one_line(x)}" for k, x in v.items()) + " }" if v else "{}"
    if isinstance(v, list):
        return "[" + ", ".join(one_line(x) for x in v) + "]"
    return json.dumps(v, ensure_ascii=False)


def set_line(text, path, key, value):
    """
    Replaces the line `"key": …` inside the section at `path` (e.g. results → government) with
    `value`; adds it at the end of the section if it isn't there; removes it if value is None.
    Everything else in the file stays exactly as it was.
    """
    lines = text.replace("\r\n", "\n").split("\n")
    start, depth = 0, 0
    for name in path:
        indent = "  " * (depth + 1)
        start = next(i for i in range(start, len(lines)) if lines[i].startswith(f'{indent}"{name}": {{'))
        depth += 1
    indent = "  " * depth
    end = next(i for i in range(start + 1, len(lines)) if lines[i] == f"{indent}}}" or lines[i] == f"{indent}}},")
    inner = "  " * (depth + 1)
    found = next((i for i in range(start + 1, end) if lines[i].startswith(f'{inner}"{key}": ')), None)
    if found is not None:
        if value is None:
            del lines[found]
            # The new last line in the section must not end with a comma.
            last = end - 2
            if lines[last].endswith(","):
                lines[last] = lines[last][:-1]
        else:
            comma = "," if lines[found].rstrip().endswith(",") else ""
            lines[found] = f'{inner}"{key}": {value}{comma}'
    elif value is not None:
        last = end - 1
        if not lines[last].rstrip().endswith((",", "{")):
            lines[last] += ","
        lines.insert(end, f'{inner}"{key}": {value}')
    return "\n".join(lines)


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    file = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_XLSX
    if cmd == "export":
        export(file)
    elif cmd == "import":
        import_(file)
    else:
        print(__doc__)
